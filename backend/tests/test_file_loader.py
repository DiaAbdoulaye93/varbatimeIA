"""Uploaded file parsing: every accepted format, and refusal of dangerous files."""

import io
import json
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook

from app.services.file_loader import MAX_ROWS, FileError, load_verbatims
from tests.test_preprocess import INPUT_CSV


def xlsx(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = Document()
    document.add_heading("Enquête de satisfaction", level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- accepted formats ---


def test_demo_csv() -> None:
    loaded = load_verbatims(INPUT_CSV.read_bytes(), "demo.csv")
    assert (loaded.source_format, loaded.source_column, len(loaded.texts)) == (
        "csv",
        "verbatim",
        223,
    )


def test_csv_semicolon_bom_cp1252_and_multiline() -> None:
    data = "﻿Id;Note;Verbatim\n1;3;Très bien, merci\n2;1;\n".encode()
    assert load_verbatims(data, "a.csv").texts == ["Très bien, merci", ""]
    assert load_verbatims("verbatim\nréseau coupé\n".encode("cp1252"), "b.csv").texts == [
        "réseau coupé"
    ]
    assert load_verbatims(b'verbatim\n"ligne 1\nligne 2"\n', "c.csv").texts == ["ligne 1\nligne 2"]


def test_csv_without_verbatim_column_picks_the_free_text_column() -> None:
    data = (
        "id,note,date,ce que vous en pensez\n"
        "1,5,2026-09-01,Le réseau coupe le soir\n"
        "2,2,2026-09-02,Agence très accueillante\n"
    )
    loaded = load_verbatims(data.encode(), "enquete.csv")
    assert loaded.source_column == "ce que vous en pensez"
    assert loaded.texts == ["Le réseau coupe le soir", "Agence très accueillante"]


def test_known_column_alias() -> None:
    loaded = load_verbatims("id;Réponse\n1;bien\n".encode(), "x.csv")
    assert (loaded.source_column, loaded.texts) == ("Réponse", ["bien"])


def test_tsv_and_txt() -> None:
    assert load_verbatims(b"id\tavis\n1\tbon accueil\n", "x.tsv").texts == ["bon accueil"]
    txt = "Le réseau coupe, souvent le soir\n\nAgence top\n".encode()
    loaded = load_verbatims(txt, "reponses.txt")
    assert (loaded.source_format, loaded.texts) == (
        "txt",
        ["Le réseau coupe, souvent le soir", "Agence top"],
    )


def test_json_variants() -> None:
    assert load_verbatims(json.dumps(["a", " b "]).encode(), "x.json").texts == ["a", "b"]
    objects = [{"id": 1, "commentaire": "réseau lent"}, {"id": 2, "commentaire": "ok"}]
    loaded = load_verbatims(json.dumps({"reponses": objects}).encode(), "x.json")
    assert (loaded.source_column, loaded.texts) == ("commentaire", ["réseau lent", "ok"])


def test_xlsx() -> None:
    data = xlsx(
        [["id", "Verbatim", "note"], [1, "Le réseau coupe", 2], [2, "Très bien", 5], [3, None, 4]]
    )
    loaded = load_verbatims(data, "enquete.xlsx")
    assert (loaded.source_format, loaded.source_column) == ("xlsx", "Verbatim")
    assert loaded.texts == ["Le réseau coupe", "Très bien", ""]


def test_docx_paragraphs_skip_headings() -> None:
    loaded = load_verbatims(docx(["Le réseau coupe", "", "Agence top"]), "enquete.docx")
    assert (loaded.source_format, loaded.texts) == ("docx", ["Le réseau coupe", "Agence top"])


def test_docx_table() -> None:
    data = docx(
        ["Tableau des réponses :"], table=[["N°", "Avis"], ["1", "Réseau lent"], ["2", "Bien"]]
    )
    loaded = load_verbatims(data, "enquete.docx")
    assert (loaded.source_column, loaded.texts) == ("Avis", ["Réseau lent", "Bien"])


def test_format_is_read_from_content_not_extension() -> None:
    data = xlsx([["avis"], ["bien"]])
    assert load_verbatims(data, "renamed.csv").source_format == "xlsx"


# --- refused files ---


@pytest.mark.parametrize(
    ("data", "filename", "message"),
    [
        (b"alert('x')", "payload.js", r"\.js ne sont pas acceptés"),
        (b"<html><script>x</script></html>", "page.html", r"\.html ne sont pas acceptés"),
        (b"MZ\x90\x00\x03", "virus.csv", "programme Windows"),
        (b"\x7fELF\x02\x01", "data.txt", "programme Linux"),
        (b"#!/bin/sh\nrm -rf /", "run.csv", "un script"),
        (b'<svg onload="alert(1)"></svg>', "image.csv", "balisage"),
        (b"\xd0\xcf\x11\xe0\xa1\xb1", "old.xls", "macros"),
        (b"", "vide.csv", "vide"),
        (b"\x00\x01\x02binary", "data.csv", "pas un fichier texte"),
    ],
)
def test_dangerous_or_unreadable_files_are_refused(
    data: bytes, filename: str, message: str
) -> None:
    with pytest.raises(FileError, match=message):
        load_verbatims(data, filename)


def test_macro_enabled_office_file_is_refused() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "<w:document/>")
        archive.writestr("word/vbaProject.bin", b"\x00" * 10)
    with pytest.raises(FileError, match="macros"):
        load_verbatims(buffer.getvalue(), "rapport.docx")


def test_zip_bomb_is_refused() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", b"0" * (60 * 1024 * 1024))
    with pytest.raises(FileError, match="décompressé"):
        load_verbatims(buffer.getvalue(), "bomb.xlsx")


def test_too_many_rows() -> None:
    with pytest.raises(FileError, match="Trop de réponses"):
        load_verbatims(("verbatim\n" + "x\n" * (MAX_ROWS + 1)).encode(), "big.csv")


# --- free text split into responses (numbered lists, exports of a table) ---


def test_numbered_export_with_preamble_and_wrapped_answers() -> None:
    lines = [
        "Enquete de satisfaction",
        "300 reponses - export du 22 octobre",
        "N",
        "Verbatim",
        "1",
        "Le reseau coupe",
        "le soir a Pikine",
        "2",
        "Agence top",
        "Page 2",
        "3",
        "Frais eleves",
    ]
    assert load_verbatims("\n".join(lines).encode(), "export.txt").texts == [
        "Le reseau coupe le soir a Pikine",
        "Agence top",
        "Frais eleves",
    ]


def test_numbered_txt_with_inline_numbers() -> None:
    lines = ["Enquete 2026, 3 questions", "1. Le reseau coupe", "2) Agence top", "3 - Frais eleves"]
    assert load_verbatims("\n".join(lines).encode(), "r.txt").texts == [
        "Le reseau coupe",
        "Agence top",
        "Frais eleves",
    ]


def test_numbered_word_paragraphs() -> None:
    data = docx(["1", "Le réseau coupe", "le soir", "2", "Agence top", "3", "Frais élevés"])
    assert load_verbatims(data, "export.docx").texts == [
        "Le réseau coupe le soir",
        "Agence top",
        "Frais élevés",
    ]


def test_pdf_is_refused_even_when_renamed() -> None:
    with pytest.raises(FileError, match=r"\.pdf ne sont pas acceptés"):
        load_verbatims(b"%PDF-1.4 ...", "enquete.pdf")
    with pytest.raises(FileError, match="document PDF"):
        load_verbatims(b"%PDF-1.4 ...", "enquete.csv")
